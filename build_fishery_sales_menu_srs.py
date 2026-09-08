import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import os

doc = docx.Document()

# Margins
for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

def style_table_header(row, col_widths, bg_color="1F4E79"):
    for idx, cell in enumerate(row.cells):
        cell.width = Inches(col_widths[idx])
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_color}"/>')
        cell._tc.get_or_add_tcPr().append(shading)
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in paragraph.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(9.5)
                run.font.name = 'Calibri'

def style_table_row(row, col_widths, is_even=False):
    bg_color = "F2F4F7" if is_even else "FFFFFF"
    for idx, cell in enumerate(row.cells):
        cell.width = Inches(col_widths[idx])
        if is_even:
            shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_color}"/>')
            cell._tc.get_or_add_tcPr().append(shading)
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(9)
                run.font.name = 'Calibri'

def add_callout(doc, title, text, bg_color="EBF3FA", border_color="2563EB"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_color}"/>')
    tcPr.append(shd)
    
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="36" w:space="0" w:color="{border_color}"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    r_title = p.add_run(f"📌 {title}\n")
    r_title.bold = True
    r_title.font.name = 'Calibri'
    r_title.font.size = Pt(10)
    r_title.font.color.rgb = RGBColor(31, 78, 121)
    
    r_text = p.add_run(text)
    r_text.font.name = 'Calibri'
    r_text.font.size = Pt(9.5)
    r_text.font.color.rgb = RGBColor(51, 51, 51)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

# =========================================================================
# TITLE SECTION
# =========================================================================
title_p = doc.add_paragraph()
title_p.paragraph_format.space_before = Pt(20)
title_p.paragraph_format.space_after = Pt(4)
r_main_title = title_p.add_run("FISHERY SALES APP — COMPLETE MENU & SUBMENU SRS")
r_main_title.font.name = 'Calibri'
r_main_title.font.size = Pt(22)
r_main_title.font.bold = True
r_main_title.font.color.rgb = RGBColor(31, 78, 121)

sub_p = doc.add_paragraph()
sub_p.paragraph_format.space_after = Pt(16)
r_sub = sub_p.add_run("Granular Functional Specification of All Navigation Menus, Screens, User Actions, and Workflows")
r_sub.font.name = 'Calibri'
r_sub.font.size = Pt(12)
r_sub.font.italic = True
r_sub.font.color.rgb = RGBColor(89, 89, 89)

doc.add_page_break()

# =========================================================================
# MENUS & SUBMENUS DETAILED SECTIONS
# =========================================================================

menu_sections = [
    {
        "menu": "MENU 1: DASHBOARD & EXECUTIVE OVERVIEW",
        "submenus": [
            ("1.1 Live Depot Floor Overview", 
             "Aim: Real-time visual radar of ongoing operations on the depot floor.\n"
             "Work: Displays real-time KPI tiles: Total Fish Arrived Today (Kg), Crates Sorted, Spot Sales Completed (₹), Boxes Ready for Dispatch, Crates in Cold Storage, and Rejection Loss % (Rotten).\n"
             "Inputs: Date selector, Depot selector, Auto-refresh toggle (15s / 30s).\n"
             "Outputs: Interactive live floor counters and quick shortcut buttons to Weighing Scale, Spot POS, and Dispatch creation."),
            
            ("1.2 Intake & Quality Yield Analytics",
             "Aim: High-level analytics of incoming quality distributions.\n"
             "Work: Charts comparing Fresh vs Rotten vs Spoiled percentages per catching point, identifying high-spoilage transit points.\n"
             "Inputs: Date range, Point filter, Species filter.\n"
             "Outputs: Species breakdown pie charts and point-wise quality yield bar graphs."),
            
            ("1.3 Intercity Dispatch Tracker",
             "Aim: Monitor outbound wholesale truck shipments across destination cities.\n"
             "Work: Live map and shipment cards showing vehicles in transit, destination cities (Delhi, Kolkata, Mumbai), departure time, estimated arrival, and total metric tons dispatched.\n"
             "Inputs: Status filter (Loading, Dispatched, Delivered, Delayed).\n"
             "Outputs: Fleet shipment cards with driver direct call button and delivery acknowledgment status.")
        ]
    },
    {
        "menu": "MENU 2: MASTER SETUP & CONFIGURATION",
        "submenus": [
            ("2.1 Catching Points & Ghats Master",
             "Aim: Manage all peripheral reservoir locations where fish are harvested.\n"
             "Work: Registers Point Code, Point Name (e.g. Ghat A, Ghat B, Deolond, Markandeya), Transport Mode (Boat, Tractor, Truck), Distance to Depot (Km), and Supervisor In-charge.\n"
             "Inputs: Point Name, Code, Geolocation, Transport Type, Contact Person.\n"
             "Outputs: Active list of points available in Intake Arrival and Toll screens."),
            
            ("2.2 Processing Depots Master",
             "Aim: Configure central processing and packaging depots.\n"
             "Work: Registers Central Depots, cold storage room capacities, digital weighing scale device IDs, and thermal printer IP/Bluetooth endpoints.\n"
             "Inputs: Depot Name, Address, Cold Room Capacity (MT), Hardware Peripherals.\n"
             "Outputs: Depot configuration profile."),
            
            ("2.3 Destination Cities & Markets",
             "Aim: Configure intercity wholesale selling destinations.\n"
             "Work: Registers destination cities (Delhi Gazipur, Kolkata Howrah, Mumbai Sassoon Dock, Nagpur), standard route distance (Km), transit duration, and default commission agents.\n"
             "Inputs: City Name, Wholesale Market Name, Distance (Km), Primary Buyer Contact.\n"
             "Outputs: Destination dropdown in Dispatch Builder."),
            
            ("2.4 Species & 3-Tier Quality Catalog",
             "Aim: Maintain species variants and quality grading pricing standards.\n"
             "Work: Configures fish species (Rohu, Catla, Mrigal, Kalbasu, Aor, Lanchi, Local Minor, Sawal) and sets rules for the 3 Quality Tiers: 1. Fresh (Standard commercial price), 2. Rotten (Rejected/0 rate), 3. Spoiled/Damaged (Discounted processing price).\n"
             "Inputs: Species Name, Local Name, Tier Multipliers, Minimum Size (Kg).\n"
             "Outputs: Graded catalog for Toll Matrix and Spot POS Invoicing."),
            
            ("2.5 Packaging & Standard Tare Master",
             "Aim: Define standard container specifications.\n"
             "Work: Configures standard crate/carret capacity (35-40 Kg), crate tare weight (2.5 Kg), dispatch box dimensions, and crushed ice allowance percentage.\n"
             "Inputs: Package Type (Carret / Box), Capacity (Kg), Tare Weight (Kg), Ice %.\n"
             "Outputs: Automatic net weight calculation on digital scales."),
            
            ("2.6 Customer & Buyer Directory",
             "Aim: Maintain client master for spot traders and intercity wholesale buyers.\n"
             "Work: Registers Buyer Name, Firm Name, Mobile, GSTIN, Address, Buyer Category (Local Spot Buyer vs Intercity Wholesale Trader), Credit Limit (₹), and Payment Terms.\n"
             "Inputs: KYC, Phone, Bank Account, Credit Period (Days).\n"
             "Outputs: Buyer profile linked to POS and Wholesale Invoicing."),
            
            ("2.7 Fleet & Transport Vehicles",
             "Aim: Register all logistics transport vehicles.\n"
             "Work: Registers Vehicle Number, Vehicle Type (Insulated Reefer Truck, Open Pickup, Tractor), Carrying Capacity (Boxes / Tonnes), GPS Device ID, and Road Permit Expiry.\n"
             "Inputs: RC Number, Vehicle Model, Capacity, Insurance/Permit Dates.\n"
             "Outputs: Vehicle fleet manifest for Dispatch shipments."),
            
            ("2.8 Drivers Directory",
             "Aim: Maintain authorized drivers roster.\n"
             "Work: Registers Driver Name, Mobile Number, Driving License Number, License Expiry, and Emergency Contact.\n"
             "Inputs: Driver KYC, DL Document, Mobile Number.\n"
             "Outputs: Driver selection in Dispatch Trip Sheet.")
        ]
    },
    {
        "menu": "MENU 3: DEPOT INTAKE, QUALITY SORTING & SMART TOLL",
        "submenus": [
            ("3.1 Point Inward Arrival Manifest",
             "Aim: Log incoming transport vehicles and crates arriving from points.\n"
             "Work: Gate entry operator logs Point Name, Vehicle Type (Boat/Tractor/Truck), Vehicle Number, Estimated Crate Count, and Arrival Timestamp.\n"
             "Inputs: Point selection, Vehicle details, Inward Crate Count.\n"
             "Outputs: Inward Gate Pass and Staging Bay assignment."),
            
            ("3.2 3-Tier Quality Sorting Screen",
             "Aim: Floor labor quality grading before weighing.\n"
             "Work: Labor separates arriving fish into 3 piles: 1. Fresh (Prime), 2. Rotten (Scrap/Disposed), 3. Spoiled/Damaged. Operator verifies batch separation.\n"
             "Inputs: Inward Batch ID, Quality Grade Tagging.\n"
             "Outputs: Graded lots ready for the Weighing Scale."),
            
            ("3.3 Live Prioritized Toll Screen (Priority Rules A to E)",
             "Aim: High-speed prioritized digital weighing per point.\n"
             "Work: Implements strict priority rules:\n"
             "  - Priority A: Single-species carrets (e.g. pure Rohu 35-40 kg).\n"
             "  - Priority B: Avoid multi-species mixing when single volume is available.\n"
             "  - Priority C: Max 2 variants (e.g. Rohu + Catla mix 35-40 kg).\n"
             "  - Priority D: Prevent 3+ species in one crate.\n"
             "  - Priority E: Odd/small leftovers (<35 kg) set aside to staging bay.\n"
             "Inputs: Point ID, Quality Grade, Species Variant(s), Digital Scale Weight (Kg), Piece Count.\n"
             "Outputs: Printed Toll Ticket and Live Depot Production Increment."),
            
            ("3.4 Bachat Toll (Mix Point Leftover Toll)",
             "Aim: Reconcile odd leftover crates across all points.\n"
             "Work: After all individual point tolls are complete, all staged leftover crates from all points are pooled and weighed as 'Bachat Toll'.\n"
             "Inputs: Selection of staged leftover crates, Weighing scale trigger.\n"
             "Outputs: Reconciled Bachat Toll Ticket and zero floor balance for incoming points."),
            
            ("3.5 Intake History & Weighing Slips Archive",
             "Aim: Audit historical intake weighing records.\n"
             "Work: Searchable archive of all completed toll tickets by Date, Point, Operator, and Quality Grade with re-print capabilities.\n"
             "Inputs: Date range, Point filter, Ticket # search.\n"
             "Outputs: Detailed weighing logs and PDF export.")
        ]
    },
    {
        "menu": "MENU 4: 1-CLICK POS & BOX PACKAGING",
        "submenus": [
            ("4.1 1-Click Scale POS & Packaging Screen",
             "Aim: Unified high-speed floor screen executing Weighing + Spot Sale + Box Packaging in ONE click.\n"
             "Work: As the crate sits on the scale:\n"
             "  1. Live weight is automatically read from electronic scale.\n"
             "  2. If a buyer is standing at the depot, select Customer from dropdown -> automatically generates Sales Invoice in background.\n"
             "  3. If packing for dispatch, check '[x] Create Box' -> automatically assigns Box Number and prints Thermal Barcode Slip.\n"
             "Inputs: Customer (Optional), Species Variant, Scale Weight, Checkbox [x] Create Box.\n"
             "Outputs: Background Invoice (if sold) OR Thermal Label Slip (if boxed) + Recorded Production."),
            
            ("4.2 Spot Sales Counter (Live Invoicing Desk)",
             "Aim: Manage counter cash/UPI sales to local traders.\n"
             "Work: View and manage all spot invoices generated from the weighing scale, collect cash or UPI payments, and print thermal receipt.\n"
             "Inputs: Invoice #, Payment Mode (Cash, UPI, Credit), Amount Collected.\n"
             "Outputs: Payment Receipt and Cash Register credit."),
            
            ("4.3 Thermal Box Label Printing & Reprint",
             "Aim: Manage box barcode label printer queue.\n"
             "Work: Prints standardized adhesive labels: Box #, Origin Depot, Species Variant, Quality Grade, Piece Count, Net Weight (Kg), Tare Weight, Date/Time.\n"
             "Inputs: Box ID / QR scan, Reprint command.\n"
             "Outputs: 4x2 inch or 4x6 inch thermal adhesive box label."),
            
            ("4.4 Ready Box Packaging Inventory",
             "Aim: Live inventory of sealed boxes waiting for truck loading.\n"
             "Work: Lists all prepared boxes on the packaging floor grouped by species and grade, ready to be selected for Intercity Dispatch.\n"
             "Inputs: Filter by Species, Grade, Packing Date.\n"
             "Outputs: Manifest of available boxes with total weight and piece count.")
        ]
    },
    {
        "menu": "MENU 5: ICE PRESERVATION & COLD STORAGE",
        "submenus": [
            ("5.1 Floor Sweep & Closing Production Toll",
             "Aim: End-of-shift floor clearance and unsold fish recording.\n"
             "Work: After point tolls and spot sales are completed, floor workers gather all remaining loose fish, sort them, and execute a 'Closing Production Toll'.\n"
             "Inputs: Closing Toll Weight (Kg), Species composition, Floor supervisor sign-off.\n"
             "Outputs: Closing Production Record preventing unaccounted inventory shrinkage."),
            
            ("5.2 Crushed Ice Crates Sealing",
             "Aim: Log ice preservation packing for overnight storage.\n"
             "Work: Records crates packed with crushed ice, ice-to-fish weight ratio, and sealing tag numbers.\n"
             "Inputs: Crate count, Ice weight (Kg), Seal IDs.\n"
             "Outputs: Preserved crate batch ready for cold room transfer."),
            
            ("5.3 Cold Room Bay Management",
             "Aim: Track crate locations inside cold storage.\n"
             "Work: Assigns ice-sealed crates to Cold Room Bays/Racks (e.g. Bay A1, Bay B2) with temperature logging and expiry timers.\n"
             "Inputs: Bay selection, Crate batch ID, Temperature (°C).\n"
             "Outputs: Cold storage live inventory map."),
            
            ("5.4 Daily Depot Mass Balance Reconciliation",
             "Aim: Mathematical verification of daily depot fish flow.\n"
             "Work: Validates: Inward Catch = Spot Sales + Boxed Dispatches + Cold Storage Stock + Rotten Disposals. Flags discrepancies > 0.5%.\n"
             "Inputs: Auto-computed from daily modules.\n"
             "Outputs: Depot Clearance Certificate.")
        ]
    },
    {
        "menu": "MENU 6: INTERCITY DISPATCH & FLEET LOGISTICS",
        "submenus": [
            ("6.1 Create New Dispatch (Shipment Builder)",
             "Aim: Build wholesale multi-box truck consignments.\n"
             "Work: Selects Origin Depot, Destination City/Market, Consignee Buyer, and adds prepared boxes via barcode scanner or batch selection.\n"
             "Inputs: Destination City, Consignee, Box selection.\n"
             "Outputs: Consolidated Dispatch Manifest with total boxes, species breakdown, and gross weight."),
            
            ("6.2 Vehicle & Driver Trip Sheet Assignment",
             "Aim: Link vehicle fleet and drivers to dispatch shipments.\n"
             "Work: Assigns Vehicle Number, Driver Name, Mobile, Distance in Km, Departure Time, and Advance Freight Cash.\n"
             "Inputs: Vehicle dropdown, Driver dropdown, Trip Advance (₹).\n"
             "Outputs: Driver Trip Sheet and Expense Voucher."),
            
            ("6.3 Barcode Loading Scanner & Manifest Verification",
             "Aim: Verify physical box loading onto the truck.\n"
             "Work: Floor scanner beeps as each box is loaded into the truck, ensuring no wrong box or short-shipment occurs.\n"
             "Inputs: Handheld barcode / QR scanner.\n"
             "Outputs: 100% Verified Loading Confirmation."),
            
            ("6.4 Gate Pass & Delivery Challan Generator",
             "Aim: Generate legal transit documents for highway transport.\n"
             "Work: Prints official Dispatch Delivery Challan, Gate Pass, and Consignment Note for interstate highway checkpoints.\n"
             "Inputs: Finalized Dispatch ID.\n"
             "Outputs: Printed 3-copy Delivery Challan & Gate Pass."),
            
            ("6.5 In-Transit Live Tracking & Delivery Acknowledgment",
             "Aim: Track shipment transit and record destination delivery.\n"
             "Work: Monitors vehicle en route, logs destination arrival, records receiving buyer signature, destination scale weight, and transit shrinkage.\n"
             "Inputs: Delivery Timestamp, Received Weight, Buyer Signature / OTP.\n"
             "Outputs: Delivery Acknowledgment Receipt (POD) and invoice trigger.")
        ]
    },
    {
        "menu": "MENU 7: COMMERCIAL BILLING & CUSTOMER ACCOUNTS",
        "submenus": [
            ("7.1 Spot POS Invoices & Counter Cash Register",
             "Aim: Manage all counter sales invoices and cash drawer balances.\n"
             "Work: Lists all spot invoices generated during scale weighing, reconciles cash drawer with UPI settlements, and generates daily cash closure.\n"
             "Inputs: Cash drawer count, POS invoice reconciliation.\n"
             "Outputs: Daily Cash Closing Statement."),
            
            ("7.2 Wholesale Consignment Invoices",
             "Aim: Commercial billing for intercity truck dispatches.\n"
             "Work: Automatically generates wholesale commercial invoices for destination buyers based on manifest box weights, agreed city rate/kg, and freight terms.\n"
             "Inputs: Dispatch ID, Agreed Rate/Kg, Freight Terms (Paid/To-Pay).\n"
             "Outputs: Commercial Wholesale Invoice & Tax Invoice."),
            
            ("7.3 Customer Credit Ledgers & Account Statements",
             "Aim: Manage buyer credit accounts, debit balances, and aging.\n"
             "Work: Itemized customer ledger showing Opening Balance, Invoices Billed, Payments Received, Credit Notes (Transit claims), and Current Outstanding.\n"
             "Inputs: Customer filter, Date range.\n"
             "Outputs: Customer Statement of Account (PDF/Excel) and Aging Analysis."),
            
            ("7.4 Payment Receipts (NEFT / RTGS / Cash / Cheque)",
             "Aim: Record incoming customer remittances.\n"
             "Work: Records payments received against outstanding invoices with Bank Reference / UTR numbers.\n"
             "Inputs: Customer ID, Payment Mode, Amount, UTR/Ref #, Bank Name.\n"
             "Outputs: Official Payment Receipt Voucher and Ledger Credit."),
            
            ("7.5 Transit Damage & Weight Shrinkage Claims",
             "Aim: Process destination weight loss claims and credit notes.\n"
             "Work: Reconciles origin dispatched weight vs destination received weight, calculates allowable melting/shrinkage % (e.g. 1-2%), and issues Credit Notes for approved claims.\n"
             "Inputs: Dispatch ID, Destination Scale Weight, Approved Claim (Kg/₹).\n"
             "Outputs: Credit Note Voucher adjusting buyer ledger.")
        ]
    },
    {
        "menu": "MENU 8: EXECUTIVE INTELLIGENCE, AUDIT & REPORTS",
        "submenus": [
            ("8.1 Point Intake & Quality Loss Yield Report",
             "Aim: Audit point-wise catch volumes and grading rejections.\n"
             "Work: Detailed report comparing total catch per point with breakdown of Fresh %, Rotten %, and Spoiled %.\n"
             "Inputs: Date range, Point filter, Species filter.\n"
             "Outputs: Point quality rating matrix and high-spoilage alert flags."),
            
            ("8.2 Daily 3-Way Production Reconciliation Report",
             "Aim: Mathematical verification of depot fish balance.\n"
             "Work: Reconciles: Intake Weight = Spot Sales Weight + Dispatched Weight + Cold Storage Weight + Scrap Disposal.\n"
             "Inputs: Date selector, Depot selector.\n"
             "Outputs: Master Production Reconciliation Register."),
            
            ("8.3 City-wise Sales Realization & Profitability Report",
             "Aim: Evaluate financial yield across destination markets.\n"
             "Work: Compares revenue per kg across Delhi, Kolkata, Mumbai, and local counter, factoring in ice, freight, and commission costs to find most profitable markets.\n"
             "Inputs: Date range, City filter, Species filter.\n"
             "Outputs: Net Profit Realization Table per Kg."),
            
            ("8.4 Box Traceability & Lifecycle Audit Trail",
             "Aim: 100% end-to-end traceability of packaged boxes.\n"
             "Work: Enter any Box # (or scan barcode) to see: Origin Point, Harvesting Date, Graded Quality, Pack Time, Dispatch Vehicle #, Driver, Destination City, and Consignee.\n"
             "Inputs: Box Serial Number / Barcode Scan.\n"
             "Outputs: Complete visual lifecycle timeline."),
            
            ("8.5 Fleet Transit & Driver Performance Report",
             "Aim: Evaluate transport logistics efficiency.\n"
             "Work: Analyzes trip turnaround times, distance (Km), diesel expenses, transit delay hours, and destination shrinkage percentages per vehicle and driver.\n"
             "Inputs: Date range, Vehicle filter, Driver filter.\n"
             "Outputs: Logistics Efficiency Index.")
        ]
    },
    {
        "menu": "MENU 9: SYSTEM ADMINISTRATION & SECURITY",
        "submenus": [
            ("9.1 User Management & Role Permissions (RBAC)",
             "Aim: Secure user access based on organizational role.\n"
             "Work: Configures users and assigns granular permissions (Super Admin, Weighing Operator, POS Cashier, Floor Supervisor, Dispatch In-charge, Accountant, Auditor).\n"
             "Inputs: User details, Role assignment, Allowed Depots/Stations.\n"
             "Outputs: Active user credentials and permission matrix."),
            
            ("9.2 Hardware Integration & Device Settings",
             "Aim: Configure floor peripherals (Scales & Printers).\n"
             "Work: Manages COM port / Bluetooth / IP settings for digital weighing scales and thermal barcode label printers.\n"
             "Inputs: Scale Baud Rate / IP, Printer Model, Label Size (4x2 / 4x6).\n"
             "Outputs: Hardware connectivity test and active peripheral status."),
            
            ("9.3 Period-End Financial Audit Lock",
             "Aim: Enforce immutable accounting books.\n"
             "Work: Sets a cutoff date (e.g. 31/03/2026). Strictly blocks editing or deletion of toll tickets, inward entries, dispatch challans, or invoices prior to cutoff.\n"
             "Inputs: Cutoff Lock Date, Admin PIN / Password.\n"
             "Outputs: System-wide audit lock enforcement."),
            
            ("9.4 Audit Trails & Soft-Delete Trash Recovery",
             "Aim: Total forensic audit trail and data safety net.\n"
             "Work: Logs all system actions (User, Timestamp, IP, Old Value, New Value) and allows authorized admins to restore accidentally soft-deleted records.\n"
             "Inputs: Search by Table, Record ID, User.\n"
             "Outputs: Comprehensive action log and 1-click restore.")
        ]
    }
]

# Write all menus and submenus to document
for m_idx, section in enumerate(menu_sections):
    h = doc.add_heading(section["menu"], level=1)
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(6)
    
    for sm_title, sm_body in section["submenus"]:
        p_t = doc.add_paragraph()
        p_t.paragraph_format.space_before = Pt(6)
        p_t.paragraph_format.space_after = Pt(2)
        r_t = p_t.add_run(f"📂 {sm_title}")
        r_t.bold = True
        r_t.font.name = 'Calibri'
        r_t.font.size = Pt(11)
        r_t.font.color.rgb = RGBColor(31, 78, 121)
        
        p_b = doc.add_paragraph()
        p_b.paragraph_format.space_after = Pt(4)
        p_b.paragraph_format.left_indent = Inches(0.2)
        r_b = p_b.add_run(sm_body)
        r_b.font.name = 'Calibri'
        r_b.font.size = Pt(9.5)
        
    doc.add_page_break()

# Save Word Documents
target_menu_word1 = r'D:\bansagar\Fishery_Sales_App_Menu_Submenu_SRS.docx'
target_menu_word2 = r'C:\Users\akdln\.gemini\antigravity\brain\fafabb32-550e-4cba-9d4d-8f4e28e8b306\Fishery_Sales_App_Menu_Submenu_SRS.docx'

doc.save(target_menu_word1)
doc.save(target_menu_word2)

print("SUCCESS: Menu & Submenu SRS Word Document generated at:")
print("1. " + target_menu_word1)
print("2. " + target_menu_word2)
