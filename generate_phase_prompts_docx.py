import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import os

doc = docx.Document()

# Page Margins
for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

def style_table_header(row, col_widths, bg_color="1F4E79"):
    for idx, cell in enumerate(row.cells):
        cell.width = Inches(col_widths[idx])
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_color}"/>')
        cell._tc.get_or_add_tcPr().append(shading)
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in paragraph.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(9.5)
                run.font.name = 'Calibri'

def style_table_row(row, col_widths, is_even=False):
    bg_color = "F2F4F7" if is_even else "FFFFFF"
    for idx, cell in enumerate(row.cells):
        cell.width = Inches(col_widths[idx])
        if is_even:
            shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_color}"/>')
            cell._tc.get_or_add_tcPr().append(shading)
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(9)
                run.font.name = 'Calibri'

def add_callout(doc, title, text, bg_color="EBF3FA", border_color="2563EB"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_color}"/>')
    tcPr.append(shd)
    
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="36" w:space="0" w:color="{border_color}"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    r_title = p.add_run(f"📌 {title}\n")
    r_title.bold = True
    r_title.font.name = 'Calibri'
    r_title.font.size = Pt(10)
    r_title.font.color.rgb = RGBColor(31, 78, 121)
    
    r_text = p.add_run(text)
    r_text.font.name = 'Calibri'
    r_text.font.size = Pt(9.5)
    r_text.font.color.rgb = RGBColor(51, 51, 51)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_prompt_box(doc, phase_title, prompt_text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F8FAFC"/>')
    tcPr.append(shd)
    
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="8" w:space="0" w:color="CBD5E1"/>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="0F172A"/>'
        f'<w:bottom w:val="single" w:sz="8" w:space="0" w:color="CBD5E1"/>'
        f'<w:right w:val="single" w:sz="8" w:space="0" w:color="CBD5E1"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(4)
    r_title = p.add_run(f"🤖 READY-TO-RUN EXECUTION PROMPT: {phase_title}\n")
    r_title.bold = True
    r_title.font.name = 'Consolas'
    r_title.font.size = Pt(10)
    r_title.font.color.rgb = RGBColor(15, 23, 42)
    
    r_text = p.add_run(prompt_text)
    r_text.font.name = 'Consolas'
    r_text.font.size = Pt(8.5)
    r_text.font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def add_heading_1(text):
    h = doc.add_heading(text, level=1)
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(4)
    for r in h.runs:
        r.font.name = 'Calibri'
        r.font.color.rgb = RGBColor(31, 78, 121)

def add_heading_2(text):
    h = doc.add_heading(text, level=2)
    h.paragraph_format.space_before = Pt(10)
    h.paragraph_format.space_after = Pt(3)
    for r in h.runs:
        r.font.name = 'Calibri'
        r.font.color.rgb = RGBColor(59, 130, 246)

def add_bullet(text, bold_prefix=""):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(2)
    if bold_prefix:
        r_b = p.add_run(bold_prefix)
        r_b.bold = True
        r_b.font.name = 'Calibri'
        r_b.font.size = Pt(9.5)
    r_t = p.add_run(text)
    r_t.font.name = 'Calibri'
    r_t.font.size = Pt(9.5)

# =========================================================================
# COVER TITLE
# =========================================================================
title_p = doc.add_paragraph()
title_p.paragraph_format.space_before = Pt(24)
title_p.paragraph_format.space_after = Pt(4)
r_main_title = title_p.add_run("BANSAGAR RESERVOIR FISHERIES ERP")
r_main_title.font.name = 'Calibri'
r_main_title.font.size = Pt(24)
r_main_title.font.bold = True
r_main_title.font.color.rgb = RGBColor(31, 78, 121)

sub_p = doc.add_paragraph()
sub_p.paragraph_format.space_after = Pt(14)
r_sub = sub_p.add_run("Complete Phase-by-Phase AI Execution Prompts Guide (100% PHP Logic & Modern Architecture)")
r_sub.font.name = 'Calibri'
r_sub.font.size = Pt(13)
r_sub.font.italic = True
r_sub.font.color.rgb = RGBColor(89, 89, 89)

meta_tbl = doc.add_table(rows=5, cols=2)
meta_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
meta_data = [
    ("Document Purpose:", "Ready-to-execute sequential prompts for AI coding assistant / Antigravity"),
    ("Methodology:", "6-Phase Incremental Architecture (Zero Logic & Calculation Missed)"),
    ("Tech Stack:", "React + Syncfusion + Dexie.js + Socket.IO + Fastify + PostgreSQL (UUIDv7)"),
    ("Coverage:", "All 11 Modules, 13 Masters, Weighing Floor, Dispatches, Wages Payroll & Reports"),
    ("Instructions:", "Copy and paste each phase prompt sequentially into the assistant.")
]
for i, (k, v) in enumerate(meta_data):
    ck = meta_tbl.cell(i, 0)
    cv = meta_tbl.cell(i, 1)
    ck.width = Inches(1.8)
    cv.width = Inches(4.7)
    
    pk = ck.paragraphs[0]
    pk.paragraph_format.space_after = Pt(2)
    rk = pk.add_run(k)
    rk.bold = True
    rk.font.size = Pt(9.5)
    rk.font.name = 'Calibri'
    rk.font.color.rgb = RGBColor(31, 78, 121)
    
    pv = cv.paragraphs[0]
    pv.paragraph_format.space_after = Pt(2)
    rv = pv.add_run(v)
    rv.font.size = Pt(9.5)
    rv.font.name = 'Calibri'

doc.add_page_break()

# =========================================================================
# PHASE 1 PROMPT
# =========================================================================
add_heading_1("Phase 1: Foundation, RBAC, Component Library, Dexie DB & 13 Masters")
doc.add_paragraph("Execute this prompt first to scaffold the entire backend and frontend foundations, offline Dexie schemas, and all 13 Master configuration screens.")

phase_1_prompt = """PHASE 1 EXECUTION PROMPT: FOUNDATION, MASTER HOC, UI LIBRARY, DEXIE DB & 13 MASTERS

Please implement Phase 1 of the Bansagar Fisheries ERP according to the following strict technical requirements:

1. BACKEND FOUNDATION (Fastify + Prisma + PostgreSQL + UUIDv7):
   - Set up Fastify server in `backend/` with TypeScript, CORS, and Zod error handler.
   - Configure PostgreSQL with UUIDv7 generator function `generate_uuid_v7()`.
   - Implement Prisma schema for:
     * Administrator, AdministratorGroup, AdminPageMenu, GroupPermission (7 actions: canView, canAdd, canEdit, canDelete, canExport, canPrint, canLock).
     * GroupType (Govt Charges %), MainGroup (Samiti with Bank/IFSC), Fisherman (Code, Name, Major/Minor/Sawal Fees, Opening Balance, Bank/IFSC).
     * FishingPoint (Catching Ghats, distance Km, supervisor), ProcessingDepot, DestinationMarket (Delhi, Kolkata, Mumbai), Species (Rohu, Catla, Mrigal, Kalbasu, Aor, Sawal, Lanchi, Minor with default rates), PackagingTare (35-40kg crate, 2.5kg tare), Customer (Buyers), FleetVehicle (GPS ID, permit dates), Driver, Boat (Nav), and StoreGear (Jaal products).
   - Implement JWT authentication and `requirePermission(module, action)` middleware decorator.
   - Implement Socket.IO server attached to Fastify for real-time broadcasts.
   - Implement sync endpoints: `POST /api/v1/sync/push` (batch processor) & `GET /api/v1/sync/pull` (delta changes).

2. FRONTEND FOUNDATION (React + Vite + TypeScript + Syncfusion):
   - Set up React 18/19 SPA with Tailwind CSS and Syncfusion Essential JS 2 (`@syncfusion/ej2-react-*`).
   - Create Dexie.js database in `src/db/localDb.ts` matching all master tables and a `syncQueue` table for 100% offline CRUD.
   - Create Master HOC `src/hoc/withERPPage.tsx` handling:
     * RBAC permission check, global hotkeys (F1-F12, Esc, Ctrl+S), Tally-style Enter traversal (`handleEnterKeyTraversal`), offline status pill, pending outbox count, and `LiveSyncToggle` switch.
   - Create Reusable Atomic Component Library in `src/components/ui/`:
     * `ERPButton` (with permission & hotkey badge), `ERPNumericInput` (weights 0.000 & currency), `ERPTextInput`, `ERPDropdown`, `ERPAutoComplete`, `ERPDatePicker` (DD/MM/YYYY), `ERPDataGrid` (with Excel/PDF export & aggregates), `ERPDialog`, `ERPSwitch`, `ERPCard`.
   - Implement Socket.IO client hook `useRealtimeSync.ts` that updates Dexie and re-renders grids.

3. MASTER CRUD SCREENS (MENU 2):
   - Build all 13 Master pages wrapped with `withERPPage`:
     * GroupTypesPage, SamitisMainGroupPage, FishermenDirectoryPage, FishingPointsPage, ProcessingDepotsPage, DestinationMarketsPage, SpeciesCatalogPage, PackagingTarePage, CustomersBuyersPage, FleetVehiclesPage, DriversDirectoryPage, BoatsNavMasterPage, StoreGearJaalPage.
   - Every grid must support Excel export, search, pagination, and inline CRUD dialogs using custom `<ERP.../>` components.
"""
add_prompt_box(doc, "PHASE 1", phase_1_prompt)
doc.add_page_break()

# =========================================================================
# PHASE 2 PROMPT
# =========================================================================
add_heading_1("Phase 2: Live Weighing Floor, Digital Scale API, 8-Species Toll & Spot POS")
doc.add_paragraph("Execute this prompt to build the core weighing floor engine, digital scale auto-reading, and retail cash counter.")

phase_2_prompt = """PHASE 2 EXECUTION PROMPT: LIVE WEIGHING FLOOR, 8-SPECIES TOLL, SCALE API & SPOT POS

Please implement Phase 2 of the Bansagar Fisheries ERP according to the following strict operational requirements:

1. WEB SERIAL API DIGITAL SCALE INTEGRATION:
   - Create `src/services/scaleService.ts` and `src/hooks/useWeighingScale.ts` using `navigator.serial`.
   - Connect to RS-232 / USB digital scale (9600-8-N-1) and continuously parse ASCII pulses (e.g. `ST,GS,+038.500kg`).
   - Implement weight stability lock (variance <= 50g over 3 pulses) and manual keyboard Numpad entry fallback with `entry_mode: "AUTO_SCALE" | "MANUAL"`.

2. 8-SPECIES DAILY INTAKE TOLL (MENU 3):
   - Implement `DailyToll` (header: date, pointId, slipNo, addedBy, totalQty, totalWeight) and `DailyTollInfo` in Prisma & Dexie.
   - Build `PointArrivalManifestPage.tsx` (3.1) for gate inward arrival logging.
   - Build `QualitySortingPage.tsx` (3.2) for 3-tier grading (Fresh, Rotten, Spoiled).
   - Build `LiveWeighingFloorPage.tsx` (3.3):
     * Captures 8 species: Catla (Qty/Wt), Rohu (Qty/Wt), Mrigal (Qty/Wt), Kalbasu (Qty/Wt), Aor (Qty/Wt), Sawal (Qty/Wt), Lanchi (Qty/Wt), Local Minor (Wt).
     * Auto-calculates: `Total Major Qty = Cqty + Rqty + Mqty + Kqty`, `Total Major Wt = Cwt + Rwt + Mwt + Kwt`, `Total Net Wt = Gross - Tare`.
     * Implements Priority Rules A to E (Single species priority, 2-variant limit, odd crate staging).
     * 100% Keyboard-first: Numpad tare presets (Numpad 1 = 2.5kg Blue crate, Numpad 2 = 3.8kg Black tub, Numpad 0 = 0kg), Enter-key auto field advance.
   - Build `BachatTollPage.tsx` (3.4) for leftover mix pooled crate reconciliation.
   - Build `IntakeSlipsArchivePage.tsx` (3.5) with thermal slip reprint & Excel export.
   - Build `ChangeTollRatePage.tsx` (3.6) for purchase rate revisions.

3. 1-CLICK SCALE POS & BOX PACKAGING (MENU 4):
   - Build `OneClickScalePOSPage.tsx` (4.1): Scale weight triggers live Spot Invoicing OR Box Packaging with 1 click.
   - Build `SpotSalesCounterPage.tsx` (4.2): Cash drawer register, UPI settlement, and 0.1s thermal receipt print.
   - Build `ThermalLabelQueuePage.tsx` (4.3): 4x2 inch barcode printing (Box ID, Species, Grade, Net Wt, Timestamp).
   - Build `ReadyBoxInventoryPage.tsx` (4.4): Manifest of sealed boxes ready for truck dispatch.
"""
add_prompt_box(doc, "PHASE 2", phase_2_prompt)
doc.add_page_break()

# =========================================================================
# PHASE 3 PROMPT
# =========================================================================
add_heading_1("Phase 3: Cold Storage, Mass Balance & Intercity Dispatch Logistics")
doc.add_paragraph("Execute this prompt to implement fish preservation, cold storage bay tracking, and wholesale fleet shipments to Delhi, Mumbai, and Kolkata.")

phase_3_prompt = """PHASE 3 EXECUTION PROMPT: COLD STORAGE, MASS BALANCE & INTERCITY DISPATCH LOGISTICS

Please implement Phase 3 of the Bansagar Fisheries ERP according to the following strict logistics requirements:

1. ICE PRESERVATION & COLD STORAGE (MENU 5):
   - Implement `ColdStorageCrate`, `ColdRoomBay`, and `ClosingToll` in Prisma & Dexie.
   - Build `FloorSweepClosingPage.tsx` (5.1) for end-of-shift unsold loose fish recording.
   - Build `IceCratesSealingPage.tsx` (5.2) for logging crushed ice crate ratios and sealing tags.
   - Build `BayManagementPage.tsx` (5.3) for visual cold storage bay inventory and temperature logs.
   - Build `MassBalanceReconciliationPage.tsx` (5.4) automating mathematical verification:
     * `Inward Catch Wt = Spot Sales Wt + Boxed Dispatch Wt + Cold Stock Wt + Scrap/Rotten Wt`.
     * Flags alert if discrepancy > 0.5%.

2. INTERCITY DISPATCH & FLEET LOGISTICS (MENU 6):
   - Implement `DispatchConsignment`, `DispatchBoxItem`, `TripSheet`, and `DeliveryChallan` in Prisma & Dexie.
   - Build `CreateShipmentPage.tsx` (6.1) selecting Origin Depot, Destination City (Delhi, Kolkata, Mumbai), Consignee Buyer, and adding sealed boxes.
   - Build `TripSheetAssignmentPage.tsx` (6.2) assigning Reefer Truck, Driver, Distance (Km), and Advance Freight cash.
   - Build `BarcodeLoadingVerifyPage.tsx` (6.3) with handheld 2D barcode scanner integration:
     * Green 'Ding' on correct box, loud buzzer alert on wrong-truck/wrong-city mismatch.
   - Build `GatePassChallanPage.tsx` (6.4) printing official 3-copy Delivery Challan, Gate Pass, and NIC e-Way Bill JSON generator (HSN 0302).
   - Build `TransitTrackingPage.tsx` (6.5) for destination delivery acknowledgment (POD), destination weight, and transit shrinkage claims.
"""
add_prompt_box(doc, "PHASE 3", phase_3_prompt)
doc.add_page_break()

# =========================================================================
# PHASE 4 PROMPT
# =========================================================================
add_heading_1("Phase 4: Commercial Billing, Customer Ledgers & Store Stocks")
doc.add_paragraph("Execute this prompt to implement commercial invoicing for wholesale buyers and central depot gear inventory.")

phase_4_prompt = """PHASE 4 EXECUTION PROMPT: COMMERCIAL BILLING, CUSTOMER LEDGERS & STORE STOCKS

Please implement Phase 4 of the Bansagar Fisheries ERP according to the following strict accounting & store requirements:

1. COMMERCIAL BILLING & CUSTOMER ACCOUNTS (MENU 7):
   - Implement `CommercialInvoice`, `CustomerLedger`, `PaymentReceipt`, `CreditNoteClaim`, and `CashDeposit` in Prisma & Dexie.
   - Build `SpotInvoicesCashRegisterPage.tsx` (7.1) for counter invoice management and daily cash closure.
   - Build `WholesaleInvoicesPage.tsx` (7.2) generating wholesale commercial tax invoices from completed truck dispatches.
   - Build `CustomerLedgersPage.tsx` (7.3) with itemized statements (Billed, Paid, Claims, Outstanding Aging) and PDF/Excel export.
   - Build `PaymentReceiptsPage.tsx` (7.4) recording NEFT/RTGS/UPI/Cash remittances with bank UTR numbers.
   - Build `ShrinkageClaimsPage.tsx` (7.5) processing destination transit melting claims (1-2% tolerance) with Credit Notes.
   - Build `CashDepositedPage.tsx` (7.6) recording direct cash collections from fishermen and traders.

2. STORE & GEAR INVENTORY (MENU 9):
   - Implement `StoreInward`, `StoreInwardReturn`, `StoreOutward`, `StoreOutwardReturn` for nets (Jaal), twine, ice blocks, crates.
   - Build `StoreInwardPage.tsx` (9.1) & `StoreInwardReturnPage.tsx` (9.2) for supplier store purchases.
   - Build `StoreOutwardPage.tsx` (9.3) & `StoreOutwardReturnPage.tsx` (9.4) for gear issuance to fishermen/boats with ledger debiting.
"""
add_prompt_box(doc, "PHASE 4", phase_4_prompt)
doc.add_page_break()

# =========================================================================
# PHASE 5 PROMPT
# =========================================================================
add_heading_1("Phase 5: Fishermen & Samiti Payroll Engine (Wages, Advances, Bank Payout)")
doc.add_paragraph("Execute this prompt to build the complete fortnightly fishermen wages calculator matching 100% of the legacy PHP formula.")

phase_5_prompt = """PHASE 5 EXECUTION PROMPT: FISHERMEN & SAMITI PAYROLL ENGINE (WAGES & ADVANCES)

Please implement Phase 5 of the Bansagar Fisheries ERP according to the following strict payroll & liability algorithms:

1. PAYROLL & WAGES ENGINE (MENU 8):
   - Implement `Wages`, `WagesItem`, `WagesAdvance`, `TransferredLiability`, and `LiabilityDeduction` in Prisma & Dexie.
   - Build `WagesSheetPage.tsx` (8.1) executing the exact fortnightly payroll formula:
     * Catch Weight Summation: `Major Wt = SUM(Catla + Rohu + Mrigal + Kalbasu)`, `Minor Wt = SUM(Local Minor)`, `Sawal Wt = SUM(Sawal)`.
     * Gross Earnings = `(Major Wt × MajorFee) + (Minor Wt × MinorFee) + (Sawal Wt × SawalFee)`.
     * Deductions = `(Gross × Govt Charges %) + Cash Advances + Liability Recovery Deductions`.
     * `Net Payable = Gross Earnings - Total Deductions`.
   - Build `FishermanAdvancesPage.tsx` (8.2) & `BulkAdvanceWagesPage.tsx` (8.3) for seasonal cash loans.
   - Build `SamitiAdvanceWagesPage.tsx` (8.4) & `SamitiWagesSheetPage.tsx` (8.5) for cooperative society consolidated statements.
   - Build `TransferredLiabilityPage.tsx` (8.6) & `LiabilityDeductionsPage.tsx` (8.7) for shifting debts between accounts.
   - Build `BankPayoutExportPage.tsx` (8.8) generating direct bank payment files (Account No, IFSC, Beneficiary Name, Net Amount) for 1-click NEFT batch upload.
"""
add_prompt_box(doc, "PHASE 5", phase_5_prompt)
doc.add_page_break()

# =========================================================================
# PHASE 6 PROMPT
# =========================================================================
add_heading_1("Phase 6: Executive Reports, Audit System & Legacy Data Migration")
doc.add_paragraph("Execute this prompt for executive intelligence reports, audit locks, and automated MySQL to PostgreSQL migration.")

phase_6_prompt = """PHASE 6 EXECUTION PROMPT: EXECUTIVE REPORTS, AUDIT LOCK & LEGACY DATA MIGRATION

Please implement Phase 6 of the Bansagar Fisheries ERP according to the following final production requirements:

1. EXECUTIVE INTELLIGENCE & AUDIT REPORTS (MENU 10):
   - Build `ProductionYieldReportPage.tsx` (10.1), `MassProductionReconciliationReportPage.tsx` (10.2), `SamitiProductionReportPage.tsx` (10.3).
   - Build `InactiveFishermenReportPage.tsx` (10.4), `FishPercentAnalysisReportPage.tsx` (10.5), `FishermanLedgerStatementPage.tsx` (10.6).
   - Build `P1P2RegulatoryReportPage.tsx` (10.7) for Govt Fisheries Department statutory filings.
   - Build `OutwardDispatchSalesReportPage.tsx` (10.8), `OutwardLedgerReportPage.tsx` (10.9), `WagesAuditReportsPage.tsx` (10.10).
   - Build `CityProfitabilityReportPage.tsx` (10.11), `BoxTraceabilityReportPage.tsx` (10.12), `FleetPerformanceReportPage.tsx` (10.13).

2. SYSTEM ADMINISTRATION & SECURITY TOOLS (MENU 11):
   - Build `UserManagementPage.tsx` (11.1), `AdministratorRolesPage.tsx` (11.2), `PermissionMatrixPage.tsx` (11.3 - visual 7-action matrix).
   - Build `HardwareSettingsPage.tsx` (11.4 - scale baud rates & thermal printer endpoints).
   - Build `PeriodAuditLockPage.tsx` (11.5) & `ClosedDatesRegisterPage.tsx` (11.6 - permanent financial book lock).
   - Build `ApplicationSettingsPage.tsx` (11.7), `AuditLogsPage.tsx` (11.8), `TrashRecoveryPage.tsx` (11.9), `BulkActionsPage.tsx` (11.10).
   - Build `DatabaseToolsPage.tsx` (11.11 - yearly archiving & balance forwarding) & `ExcelImportCenterPage.tsx` (11.12).

3. AUTOMATED LEGACY DATA MIGRATION SCRIPT:
   - Create `backend/src/scripts/migrateLegacyData.ts` connecting to legacy MySQL `bansagar` database.
   - Convert all historical Fishermen, Samitis, Fishing Points, Species, Daily Tolls, Outward Dispatches, and Wages into PostgreSQL with UUIDv7, preserving all historical balances and relationships.
"""
add_prompt_box(doc, "PHASE 6", phase_6_prompt)

# Output Path
output_file_path = os.path.join("d:\\bansagar", "Bansagar_Fisheries_Complete_PhaseWise_Execution_Prompts.docx")
doc.save(output_file_path)
print(f"Successfully generated: {output_file_path}")
